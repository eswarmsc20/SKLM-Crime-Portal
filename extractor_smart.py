import os
import json
import datetime
import openpyxl

base_dir = r"c:\Users\seeka\Desktop\SKLM\Point-1 & 2 All Head-wise Grave & Non-Grave UI Cases-upto\Point-1 & 2 All Head-wise Grave & Non-Grave UI Cases-upto"
tc_dir = os.path.join(base_dir, "Total Cases")
out_dir = r"c:\Users\seeka\Desktop\SKLM\data"
os.makedirs(out_dir, exist_ok=True)

def format_cell(v, is_pct_col=False):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float):
        if is_pct_col:
            if -10.0 <= v <= 10.0 and abs(v) > 0.00001:
                return f"{v * 100:+.1f}%"
        if v.is_integer():
            return int(v)
        return round(v, 2)
    s = str(v).strip()
    if s.startswith("'") and s.endswith("'"):
        s = s[1:-1]
    return s

def extract_sheet_table_smart(ws, max_rows=None):
    peek_rows = []
    row_iter = ws.iter_rows(values_only=True)
    
    for _ in range(12):
        try:
            r = next(row_iter)
            peek_rows.append(r)
        except StopIteration:
            break
            
    if not peek_rows:
        return {"headers": [], "rows": [], "total_rows": 0, "shown_rows": 0}
        
    # Search for header row using indicator keywords
    header_idx = -1
    for idx, r in enumerate(peek_rows):
        line = " ".join(str(c) for c in r if c is not None).lower()
        if any(k in line for k in ["s. no", "s.no", "name of the ps", "name of the i.o", "police station", "crime head", "fir no", "gravity type", "above 2 years", "overall crime statistics"]):
            header_idx = idx
            break
            
    if header_idx == -1:
        # Fallback: first row with at least 2 non-empty values
        for idx, r in enumerate(peek_rows):
            non_empty = [c for c in r if c is not None and str(c).strip() != ""]
            if len(non_empty) >= 2:
                header_idx = idx
                break
        if header_idx == -1:
            header_idx = 0

    header_row = peek_rows[header_idx]
    
    # Check if next row is a sub-header (e.g. Row 3 has groups, Row 4 has column names)
    sub_header_row = None
    if header_idx + 1 < len(peek_rows):
        next_r = peek_rows[header_idx + 1]
        next_line = " ".join(str(c) for c in next_r if c is not None).lower()
        # If it looks like numbering '(1) (2)...' skip it
        if "(1)" in next_line and "(2)" in next_line:
            pass
        elif any(k in next_line for k in ["no. of", "police disposal", "sc/prc", "actual ui", "finalized", "reported", "jan", "feb"]):
            sub_header_row = next_r
            
    headers = []
    seen = {}
    
    current_parent = ""
    for i in range(len(header_row)):
        top_h = str(header_row[i]).strip().replace('\n', ' ') if header_row[i] is not None else ""
        sub_h = str(sub_header_row[i]).strip().replace('\n', ' ') if sub_header_row and i < len(sub_header_row) and sub_header_row[i] is not None else ""
        
        if top_h and not sub_h:
            h_str = top_h
            current_parent = top_h
        elif sub_h and top_h:
            h_str = f"{top_h} - {sub_h}"
            current_parent = top_h
        elif sub_h and not top_h:
            if current_parent and "s. no" not in current_parent.lower() and "name" not in current_parent.lower():
                h_str = f"{current_parent} - {sub_h}"
            else:
                h_str = sub_h
        else:
            h_str = f"Col_{i+1}"
            
        if not h_str:
            h_str = f"Col_{i+1}"
            
        if h_str in seen:
            seen[h_str] += 1
            h_str = f"{h_str}_{seen[h_str]}"
        else:
            seen[h_str] = 1
        headers.append(h_str)
        
    pct_cols = {i for i, h in enumerate(headers) if "%" in h.lower() or "percent" in h.lower()}
    
    start_row_idx = header_idx + (2 if sub_header_row else 1)
    
    data_rows = []
    # Collect rows from peek
    for r in peek_rows[start_row_idx:]:
        line = " ".join(str(c) for c in r if c is not None).strip()
        # Skip numbering rows like '(1) | (2) | (3)'
        if line.startswith("(1)") or line.startswith("1 | 2 | 3"):
            continue
        if not any(c is not None and str(c).strip() != "" for c in r):
            continue
        cleaned = [format_cell(r[i] if i < len(r) else None, i in pct_cols) for i in range(len(headers))]
        data_rows.append(cleaned)
        
    total_data_count = len(data_rows)
    
    # Stream the rest of the rows
    for r in row_iter:
        line = " ".join(str(c) for c in r if c is not None).strip()
        if line.startswith("(1)") or line.startswith("1 | 2 | 3"):
            continue
        total_data_count += 1
        if max_rows and len(data_rows) >= max_rows:
            continue
        if not any(c is not None and str(c).strip() != "" for c in r):
            continue
        cleaned = [format_cell(r[i] if i < len(r) else None, i in pct_cols) for i in range(len(headers))]
        data_rows.append(cleaned)
        
    return {
        "headers": headers,
        "rows": data_rows,
        "total_rows": total_data_count,
        "shown_rows": len(data_rows)
    }

def process_file(path, limits=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    res = {}
    limits = limits or {}
    for s in wb.sheetnames:
        lim = limits.get(s, None)
        res[s] = extract_sheet_table_smart(wb[s], max_rows=lim)
    wb.close()
    return res

print("Smart extractor loaded.")
