import os
import json
import datetime
import openpyxl

base_dir = r"c:\Users\seeka\Desktop\SKLM\Point-1 & 2 All Head-wise Grave & Non-Grave UI Cases-upto\Point-1 & 2 All Head-wise Grave & Non-Grave UI Cases-upto"
tc_dir = os.path.join(base_dir, "Total Cases")
out_dir = r"c:\Users\seeka\Desktop\SKLM\data"
os.makedirs(out_dir, exist_ok=True)

def format_cell(v, is_pct_col=False):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float):
        if is_pct_col:
            if -10.0 <= v <= 10.0:
                pct = v * 100
                return f"{pct:+.1f}%" if abs(pct) >= 0.05 else "0.0%"
        if v.is_integer():
            return int(v)
        return round(v, 2)
    s = str(v).strip()
    if s.startswith("'") and s.endswith("'"):
        s = s[1:-1]
    return s

def extract_sheet_table_smart(ws, max_rows=None):
    peek_rows = []
    row_iter = ws.iter_rows(values_only=True)
    
    for _ in range(12):
        try:
            r = next(row_iter)
            peek_rows.append(r)
        except StopIteration:
            break
            
    if not peek_rows:
        return {"headers": [], "rows": [], "total_rows": 0, "shown_rows": 0}
        
    header_idx = -1
    for idx, r in enumerate(peek_rows):
        line = " ".join(str(c) for c in r if c is not None).lower()
        # Prefer row containing column-specific keywords over general titles
        if any(k in line for k in ["s. no.", "s.no", "name of the ps", "name of the i.o", "police station", "fir no", "gravity type", "above 2 years"]):
            header_idx = idx
            break
        elif "crime head" in line and not line.startswith("overall crime statistics"):
            header_idx = idx
            break
            
    if header_idx == -1:
        for idx, r in enumerate(peek_rows):
            non_empty = [c for c in r if c is not None and str(c).strip() != ""]
            if len(non_empty) >= 2:
                header_idx = idx
                break
        if header_idx == -1:
            header_idx = 0

    header_row = peek_rows[header_idx]
    
    # Check if next row is a sub-header
    sub_header_row = None
    if header_idx + 1 < len(peek_rows):
        next_r = peek_rows[header_idx + 1]
        next_line = " ".join(str(c) for c in next_r if c is not None).lower()
        if "(1)" in next_line and "(2)" in next_line:
            pass
        elif any(k in next_line for k in ["no. of", "police disposal", "sc/prc", "actual ui", "finalized", "reported"]):
            sub_header_row = next_r
            
    headers = []
    seen = {}
    current_parent = ""
    for i in range(len(header_row)):
        top_h = str(header_row[i]).strip().replace('\n', ' ') if header_row[i] is not None else ""
        sub_h = str(sub_header_row[i]).strip().replace('\n', ' ') if sub_header_row and i < len(sub_header_row) and sub_header_row[i] is not None else ""
        
        if top_h and not sub_h:
            h_str = top_h
            current_parent = top_h
        elif sub_h and top_h:
            h_str = f"{top_h} - {sub_h}"
            current_parent = top_h
        elif sub_h and not top_h:
            if current_parent and "s. no" not in current_parent.lower() and "name" not in current_parent.lower():
                h_str = f"{current_parent} - {sub_h}"
            else:
                h_str = sub_h
        else:
            h_str = f"Col_{i+1}"
            
        if not h_str:
            h_str = f"Col_{i+1}"
            
        if h_str in seen:
            seen[h_str] += 1
            h_str = f"{h_str}_{seen[h_str]}"
        else:
            seen[h_str] = 1
        headers.append(h_str)
        
    pct_cols = {i for i, h in enumerate(headers) if "%" in h.lower() or "percent" in h.lower()}
    
    start_row_idx = header_idx + (2 if sub_header_row else 1)
    
    data_rows = []
    for r in peek_rows[start_row_idx:]:
        line = " ".join(str(c) for c in r if c is not None).strip()
        if line.startswith("(1)") or line.startswith("1 | 2 | 3") or "(1)" in line and "(2)" in line and len(line) < 60:
            continue
        if not any(c is not None and str(c).strip() != "" for c in r):
            continue
        cleaned = [format_cell(r[i] if i < len(r) else None, i in pct_cols) for i in range(len(headers))]
        data_rows.append(cleaned)
        
    total_data_count = len(data_rows)
    
    for r in row_iter:
        line = " ".join(str(c) for c in r if c is not None).strip()
        if line.startswith("(1)") or line.startswith("1 | 2 | 3") or "(1)" in line and "(2)" in line and len(line) < 60:
            continue
        total_data_count += 1
        if max_rows and len(data_rows) >= max_rows:
            continue
        if not any(c is not None and str(c).strip() != "" for c in r):
            continue
        cleaned = [format_cell(r[i] if i < len(r) else None, i in pct_cols) for i in range(len(headers))]
        data_rows.append(cleaned)
        
    return {
        "headers": headers,
        "rows": data_rows,
        "total_rows": total_data_count,
        "shown_rows": len(data_rows)
    }

def process_and_save(file_path, json_name, js_name, var_name, sheet_limits=None):
    fname = os.path.basename(file_path)
    print(f"Processing {fname}...", flush=True)
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheets_data = {}
    sheet_limits = sheet_limits or {}
    for s in wb.sheetnames:
        lim = sheet_limits.get(s, None)
        sheets_data[s] = extract_sheet_table_smart(wb[s], max_rows=lim)
    wb.close()
    
    # Write json
    jpath = os.path.join(out_dir, json_name)
    with open(jpath, "w", encoding="utf-8") as f:
        json.dump(sheets_data, f, ensure_ascii=False)
        
    # Write js
    jspath = os.path.join(out_dir, js_name)
    with open(jspath, "w", encoding="utf-8") as f:
        f.write(f"window.{var_name} = {json.dumps(sheets_data, ensure_ascii=False)};\n")
        
    print(f"  --> Saved {json_name} and {js_name} ({len(sheets_data)} sheets)", flush=True)
    return sheets_data

print("Starting full smart extraction...")
# 1. Master Headwise
process_and_save(
    os.path.join(base_dir, "1. Total UI Cases-Head Wise.xlsx"),
    "master_headwise.json", "master_headwise.js", "MASTER_HEADWISE",
    {"DCRB": 1500}
)

# 2. Module 1: Total Reported Cases
process_and_save(
    os.path.join(tc_dir, "1. Total Reported Cases.xlsx"),
    "module_1_reported.json", "module_1_reported.js", "MODULE_1_REPORTED",
    {"Total": 2000}
)

# 3. Module 1A: Total UI Cases
process_and_save(
    os.path.join(tc_dir, "1A. Total UI Cases.xlsx"),
    "module_1a_ui.json", "module_1a_ui.js", "MODULE_1A_UI",
    {}
)

# 4. Module 1B: UI Cases Compliance
process_and_save(
    os.path.join(tc_dir, "1B. UI Cases Compliance.xlsx"),
    "module_1b_compliance.json", "module_1b_compliance.js", "MODULE_1B_COMPLIANCE",
    {}
)

# 5. Module 2: Long Pending Reported Cases
process_and_save(
    os.path.join(tc_dir, "2. Long Pending Reported Cases.xlsx"),
    "module_2_long_reported.json", "module_2_long_reported.js", "MODULE_2_LONG_REPORTED",
    {"Total": 2000}
)

# 6. Module 2A: Long Pending UI Cases
process_and_save(
    os.path.join(tc_dir, "2A. Long Pending UI Cases.xlsx"),
    "module_2a_long_ui.json", "module_2a_long_ui.js", "MODULE_2A_LONG_UI",
    {}
)

# 7. Module 4: Head Wise Abstract
process_and_save(
    os.path.join(tc_dir, "4. Head Wise Abstract.xlsx"),
    "module_4_abstract.json", "module_4_abstract.js", "MODULE_4_ABSTRACT",
    {"Total": 2000}
)

print("\nAll 7 files processed with smart headers and percentage auto-formatting!")
